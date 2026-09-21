import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def create_executive_docx(title: str, summary_markdown: str, ico_data: dict, output_path: str):
    """
    Builds a professional executive document (.docx) using python-docx.
    """
    doc = Document()
    
    # Configure 1 inch margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Document Header
    p_meta = doc.add_paragraph()
    run_meta = p_meta.add_run("TRANSFORMAI // EXECUTIVE DELIVERABLE BRIEF")
    run_meta.font.name = "Arial"
    run_meta.font.size = Pt(9)
    run_meta.font.bold = True
    run_meta.font.color.rgb = RGBColor(255, 107, 0) # iQOO Orange

    # Main Title
    h1 = doc.add_heading(level=1)
    run_h1 = h1.add_run(title)
    run_h1.font.name = "Arial"
    run_h1.font.size = Pt(22)
    run_h1.font.bold = True
    run_h1.font.color.rgb = RGBColor(15, 23, 42)

    # Sub-metadata bar
    p_info = doc.add_paragraph()
    timestamp = ico_data.get("timestamp", "Present")
    loc = ico_data.get("location", "Edge Engine")
    obj = ico_data.get("primary_objective", "Operational alignment")
    p_info.add_run(f"Timestamp: {timestamp}  |  Location: {loc}\nObjective: {obj}\n").italic = True

    # Executive Overview
    doc.add_heading("1. Strategic Overview", level=2)
    overview_text = ico_data.get("executive_overview", "")
    p_over = doc.add_paragraph(overview_text)
    p_over.style.font.name = "Arial"

    # Key Findings
    doc.add_heading("2. Key Findings & Observations", level=2)
    findings = ico_data.get("key_findings", [])
    for f in findings:
        doc.add_paragraph(f, style='List Bullet')

    # Action Items Table
    doc.add_heading("3. Action Matrix", level=2)
    actions = ico_data.get("action_items", [])
    if actions:
        table = doc.add_table(rows=1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Owner"
        hdr_cells[1].text = "Action Item"
        hdr_cells[2].text = "Deadline"

        for item in actions:
            row_cells = table.add_row().cells
            row_cells[0].text = item.get("owner", "Team")
            row_cells[1].text = item.get("task", "")
            row_cells[2].text = item.get("deadline", "TBD")
            
    # Entities / Metrics
    metrics = ico_data.get("entities", {}).get("metrics", [])
    if metrics:
        doc.add_heading("4. Key Milestones & Metrics", level=2)
        for m in metrics:
            doc.add_paragraph(f"Milestone Metric: {m}", style='List Bullet')

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc.save(output_path)
    return output_path
